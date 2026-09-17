# -*- coding: utf-8 -*-
"""Post-process pandoc DOCX into MPCE Word template format."""
import io, sys, re, copy, zipfile, shutil
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

SRC = r'C:\Users\佬肥\PycharmProjects\PythonProject\operator-transient-stability\paper_mpce_raw.docx'
DST = r'C:\Users\佬肥\PycharmProjects\PythonProject\operator-transient-stability\paper_mpce_post.docx'

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, Emu, Inches
from lxml import etree

d = Document(SRC)
body = d.element.body
paras = d.paragraphs
texts = [p.text for p in paras]

W = qn('w:')  # wordprocessingml
A = qn('a:')  # drawingml
WP = qn('wp:')
R = qn('r:')

def find_para(prefix, start=0):
    for i, t in enumerate(texts[start:], start):
        if t.strip().startswith(prefix):
            return i, paras[i]
    raise ValueError('not found: ' + prefix)

# ---------- 1. authors: merge 4 paragraphs into one ----------
i_title = 0
# paras 1..4 are the authors
auth_paras = paras[1:5]
# build new runs: author i text + superscript 'a', comma; last with 'a,*'
# keep style of first author para
first = auth_paras[0]
# delete later author paragraphs
for p in auth_paras[1:]:
    p._p.getparent().remove(p._p)
# rebuild first author paragraph runs
# collect author names
names = [p.text.strip() for p in auth_paras]
for _r in first._p.findall('.//' + W + 'r'):
    first._p.remove(_r)
names[0] = names[0].replace('\xa0', ' ')
def add_run(p_el, text, bold=False, sup=False, italic=False):
    r = p_el.makeelement(W + 'r', {})
    if bold or sup or italic:
        rPr = r.makeelement(W + 'rPr', {})
        if bold: rPr.append(rPr.makeelement(W + 'b', {}))
        if italic: rPr.append(rPr.makeelement(W + 'i', {}))
        if sup: rPr.append(rPr.makeelement(W + 'vertAlign', {W + 'val': 'superscript'}))
        r.append(rPr)
    t = r.makeelement(W + 't', {})
    t.text = text
    r.append(t)
    p_el.append(r)
    return r

for idx, nm in enumerate(names):
    add_run(first._p, nm)
    if idx == 0:
        sup = 'a,*'
    else:
        sup = 'a'
    add_run(first._p, ' ' + sup, sup=True)
    if idx < 3:
        add_run(first._p, ', ')
# center + spacing
first.alignment = 1  # center

# ---------- affiliation paragraph after authors ----------
aff_p = first._p.makeelement(W + 'p', {})
first._p.addnext(aff_p)
from docx.text.paragraph import Paragraph
aff = Paragraph(aff_p, d.paragraphs[0]._parent)
add_run(aff_p, 'a International Energy College, Jinan University, Guangzhou, 510632, China')
aff.alignment = 1

# ---------- 2. footnote for corresponding author ----------
# footnote ref at the 'a,*' superscript run of first author (find it)
fn_ref = None
for r in first._p.findall(W + 'r'):
    t = r.find(W + 't')
    if t is not None and t.text == ' a,*':
        # add footnoteReference before the t
        fref = r.makeelement(W + 'footnoteReference', {W + 'id': '1'})
        t.addprevious(fref)
        fn_ref = r
        break
assert fn_ref is not None

# ---------- 3. abstract: delete 'Abstract' heading, prepend bold 'Abstract—' ----------
i_abs_head, p_abs_head = find_para('Abstract')
i_abs_text = i_abs_head + 1
p_abs = paras[i_abs_text]
p_abs_head._p.getparent().remove(p_abs_head._p)
# insert bold run at start
brun = p_abs._p.makeelement(W + 'r', {})
rPr = brun.makeelement(W + 'rPr', {})
rPr.append(rPr.makeelement(W + 'b', {}))
brun.append(rPr)
bt = brun.makeelement(W + 't', {})
bt.text = 'Abstract—'
brun.append(bt)
p_abs._p.insert(0, brun)

# ---------- 4. keywords -> Index Terms ----------
i_kw = None
for i, t in enumerate(texts):
    if t.strip().startswith('conformal prediction'):
        i_kw = i
        break
p_kw = paras[i_kw]
# delete all runs, rebuild
for r in p_kw._p.findall(W + 'r'):
    p_kw._p.remove(r)
add_run(p_kw._p, 'Index Terms—', bold=True)
add_run(p_kw._p, 'conformal prediction, critical clearing time, neural operator, '
                  'physics-informed machine learning, power system dynamics, transient stability.')

# ---------- 5. figure captions: add Fig. N. prefix ----------
fig_captions = [('Overview of the proposed', 'Fig. 1.'),
                ('Predicted (dashed) versus true', 'Fig. 2.'),
                ('Conformal band validity', 'Fig. 3.')]
for prefix, label in fig_captions:
    i, p = find_para(prefix)
    run = p._p.makeelement(W + 'r', {})
    t = run.makeelement(W + 't', {})
    t.text = label + '  '
    run.append(t)
    p._p.insert(0, run)

# ---------- 6. table captions -> TABLE I/II/III + caption ----------
tab_captions = [('Main results on the IEEE', 'TABLE I'),
                ('Comparison of variants and baselines', 'TABLE II'),
                ('Generalization of the multi-topology', 'TABLE III')]
for prefix, label in tab_captions:
    i, p = find_para(prefix)
    # clone the caption paragraph -> becomes the caption text below the TABLE header
    cap_p = copy.deepcopy(p._p)
    p._p.addnext(cap_p)
    # rewrite the original paragraph as 'TABLE N'
    for r in p._p.findall(W + 'r'):
        p._p.remove(r)
    add_run(p._p, label, bold=True)
    p.alignment = 1
    # caption paragraph: centered italic
    from docx.text.paragraph import Paragraph as P
    cap = P(cap_p, None)
    cap.alignment = 1
    for r in cap_p.findall(W + 'r'):
        if r.find(W + 'rPr') is None:
            rp = r.makeelement(W + 'rPr', {})
            r.insert(0, rp)
        rPr = r.find(W + 'rPr')
        if rPr.find(W + 'i') is None:
            rPr.append(rPr.makeelement(W + 'i', {}))

# ---------- 7. image sizing ----------
IMG_W_EMU = 3108960  # 3.4 in
aspects = {r'rId34': 1080/3150, r'rId49': 840/1260, r'rId60': 840/2160}
for blip in body.findall('.//' + A + 'blip'):
    rid = blip.get(R + 'embed')
    if rid not in aspects:
        continue
    ratio = aspects[rid]
    h = int(IMG_W_EMU * ratio)
    # wp:extent
    inline = blip.getparent().getparent()  # blip -> a:blip parent pic; actually walk up to wp:inline
    inline = None
    node = blip
    while node is not None and node.tag != WP + 'inline':
        node = node.getparent()
    inline = node
    if inline is None:
        continue
    ext = inline.find(WP + 'extent')
    if ext is not None:
        ext.set('cx', str(IMG_W_EMU)); ext.set('cy', str(h))
    xfrm = inline.find('.//' + A + 'xfrm')
    if xfrm is not None:
        aext = xfrm.find(A + 'ext')
        if aext is not None:
            aext.set('cx', str(IMG_W_EMU)); aext.set('cy', str(h))

# ---------- 8. table width + font ----------
for tbl in d.tables:
    tblPr = tbl._tbl.tblPr
    tblW = tblPr.find(W + 'tblW')
    if tblW is None:
        tblW = tblPr.makeelement(W + 'tblW', {})
        tblPr.append(tblW)
    tblW.set(W + 'w', '4896'); tblW.set(W + 'type', 'dxa')
    for row in tbl.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    if r.font.size is None:
                        r.font.size = Pt(9)

# ---------- 9. references: tab -> space, hanging indent ----------
for i in range(len(texts)):
    t = texts[i]
    if re.match(r'^\[\d+\]\s', t):
        p = paras[i]
        for r in p._p.findall(W + 'r'):
            for node in r.iter():
                if node.tag == W + 'tab':
                    # replace tab with space in preceding t
                    prev = node.getprevious()
                    if prev is not None and prev.tag == W + 't' and prev.text is not None:
                        prev.text = prev.text + ' '
                    r.remove(node)
        pPr = p._p.find(W + 'pPr')
        if pPr is None:
            pPr = p._p.makeelement(W + 'pPr', {})
            p._p.insert(0, pPr)
        ind = pPr.find(W + 'ind')
        if ind is None:
            ind = pPr.makeelement(W + 'ind', {})
            pPr.append(ind)
        ind.set(W + 'left', '430'); ind.set(W + 'hanging', '430')

# ---------- 10. split merged Declarations paragraph ----------
i_dec = None
for i, t in enumerate(texts):
    if t.startswith('Funding.'):
        i_dec = i
        break
p_dec = paras[i_dec]
chunks = []  # list of run-lists
cur = []
for r in p_dec._p.findall(W + 'r'):
    has_br = len(r.findall(W + 'br')) > 0
    cur.append(r)
    if has_br:
        chunks.append(cur); cur = []
if cur:
    chunks.append(cur)
# create new paragraphs after p_dec
new_ps = []
for ch in chunks[1:]:
    np = p_dec._p.makeelement(W + 'p', {})
    p_dec._p.addnext(np)
    for r in ch:
        # remove br
        for br in r.findall(W + 'br'):
            r.remove(br)
        np.append(r)
    new_ps.append(np)
# keep first chunk in original paragraph
for ch in chunks[1:]:
    pass

# ---------- 11. footnotes part ----------
# add footnotes.xml content
foot_xml = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
 '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
 '<w:footnote w:type="separator" w:id="-1"><w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr><w:r><w:separator/></w:r></w:p></w:footnote>'
 '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
 '<w:footnote w:id="1"><w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>'
 '<w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr><w:t>\u2217</w:t></w:r>'
 '<w:r><w:t xml:space="preserve">Corresponding author. Email address: 3293857014@qq.com (Zhenyu Liu)</w:t></w:r>'
 '</w:p></w:footnote></w:footnotes>')

d.save(DST)
# now patch the zip: merge footnote into existing footnotes.xml (template has one)
import zipfile as zf
my_footnote = ('<w:footnote w:id="1"><w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>'
 '<w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr><w:t>∗</w:t></w:r>'
 '<w:r><w:t xml:space="preserve">Corresponding author. Email address: 3293857014@qq.com (Zhenyu Liu)</w:t></w:r>'
 '</w:p></w:footnote>')
tmp = DST + '.tmp'
with zf.ZipFile(DST, 'r') as zin, zf.ZipFile(tmp, 'w', zf.ZIP_DEFLATED) as zout:
    for item in zin.namelist():
        data = zin.read(item)
        if item == '[Content_Types].xml':
            s = data.decode('utf-8')
            if 'footnotes+xml' not in s:
                s = s.replace('</Types>',
                    '<Override PartName="/word/footnotes.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/></Types>')
            data = s.encode('utf-8')
        elif item == 'word/_rels/document.xml.rels':
            s = data.decode('utf-8')
            if 'footnotes.xml' not in s:
                s = s.replace('</Relationships>',
                    '<Relationship Id="rIdFootnotes" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/></Relationships>')
            data = s.encode('utf-8')
        elif item == 'word/footnotes.xml':
            s = data.decode('utf-8')
            if '<w:footnote w:id="1">' not in s:
                s = s.replace('</w:footnotes>', my_footnote + '</w:footnotes>')
            data = s.encode('utf-8')
        zout.writestr(item, data)
shutil.move(tmp, DST)
print('POST-PROCESSED:', DST)
